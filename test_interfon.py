"""Videointerfonul la bloc (P7b): patru tipuri noi, o conditie, o planşa.

DE CE EXISTA. Un tip nou de element trece prin multe liste (constrangerea din baza, planşa,
editorul, lista de cantitati, memoriul, caietul) si fiecare lista uitata pierde ceva TACUT: in baza
se pierd elementele la salvare, pe planşa simbolul iese ca o doza, in memoriu aparatul nu apare.
Proba compara listele INTRE ELE si pica pe numele listei care a ramas in urma.

A doua jumatate pazeste CONDITIA planşei „Schema de distributie retea de interfon”: se aprinde
DOAR din elementele plasate (`interfon.are_interfon`), citita de numerotare si de generator — nu
dintr-un comutator trimis de cineva. Proba trece PRIN ENDPOINT, fiindca un camp scos din contract
se vede doar acolo.

Fiecare proba are SONDA ei: o modificare care TREBUIE s-o faca sa pice. O proba care nu pica pe
sonda nu dovedeste nimic.

  A. migratia contine cele patru tipuri si e aditiva (proba pe COMPORTAMENT — insert real intr-o
     tranzactie anulata — s-a facut direct pe baza, cu SQL; vezi raportul P7b)
  B. listele de inregistrare, comparate intre ele + sonda pe fiecare lista
  C. P+4: 17 planşe, IE.17 = schema de interfon, desenata; fara elemente -> 16. Plus schema de
     curenti slabi: anuntata doar cand are ce desena (altfel un bloc doar cu interfon promitea una)
  D. lista de cantitati: 1 buton de iesire per yala, 1 buton de sonerie per post
  E. sursa pe TCC, circuit separat, cu protectia EI (decizia lui Dan: MCB 10 A curba B, CYY-F 3x1,5);
     restul circuitelor TCC identice cu HEAD; un receptor oarecare ramane pe regula generala
  F. copierea P4 ia si postul interior
  G. avertismentele: apartament fara post / posturi fara panou / sistem fara sursa
  H. casa: fara butoanele noi; documentele identice cu HEAD (pixeli + randuri)
  J. citirea din baza la numerotare DOAR la bloc: casa cu baza cazuta -> numerotarea de azi, fara
     nicio citire; blocul cu baza cazuta -> eroare explicita; /api/finalize trimite `este_bloc` din
     `esteBloc` (handler-ul real). Sonde: poarta scoasa din endpoint / din ruta -> prinse
  I. tsc curat

Rulare:  python test_interfon.py
"""
import copy
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

RADACINA = os.path.dirname(os.path.abspath(__file__))
APP = os.path.join(RADACINA, "apps", "zynapse-configurator")
EDITOR = os.path.join(APP, "components", "plan-editor.tsx")
CONST = os.path.join(APP, "lib", "constants.ts")
sys.path.insert(0, RADACINA)
rele = []

import interfon  # noqa: E402

TIPURI = set(interfon.TIPURI)


def v(nume, cond, det=""):
    print("  %-70s %s %s" % (nume[:70], "OK" if cond else "**ESUAT**", det if not cond else ""))
    if not cond:
        rele.append(nume)


_HEAD = {}


def head_dir():
    """Arborele de la HEAD (git archive), extras O SINGURA DATA pe rulare. Comparatiile de
    non-regresie se fac fata de el, nu fata de discul curent."""
    if "dir" not in _HEAD:
        d = tempfile.mkdtemp(prefix="p7b_head_")
        arh = os.path.join(d, "head.tar")
        subprocess.run(["git", "archive", "--format=tar", "-o", arh, "HEAD"], cwd=RADACINA, check=True)
        rad = os.path.join(d, "head")
        os.makedirs(rad)
        subprocess.run(["tar", "-xf", arh, "-C", rad], check=True)
        _HEAD["dir"], _HEAD["tmp"] = rad, d
    return _HEAD["dir"]


# ── BLOCUL DE PROBA: P+4, trei apartamente conturate pe fiecare nivel ─────────────────────────────
NIVELURI = ["parter", "etaj", "etaj 2", "etaj 3", "etaj 4"]
# Steagurile blocului de proba din R&D (P+4 cu TCC): 16 planşe azi, ultima IE.16 = schema TCC.
# `este_bloc` il pune /api/finalize (cu `esteBloc`); fara el numerotarea nu citeste elementele deloc.
BLOC_STEAGURI = {"extra_floors": ["etaj 1", "etaj 2", "etaj 3", "etaj 4"], "has_tcc": True,
                 "has_distributie": True, "has_teg": False, "este_bloc": True}


def bloc_proba(cu_interfon=True, fara_post=(), fara_panou=False, fara_sursa=False, ap_per_nivel=3,
               niveluri=None):
    els, n = [], [0]

    def add(**k):
        n[0] += 1
        d = {"id": "el-%04d" % n[0], "project_id": "proba", "plan_type": "curenti_slabi",
             "label": None, "room": None, "wall_mounted": False, "rotation": 0, "power_w": None,
             "cable_path": [], "mount_height_m": None, "kit_panica": False}
        d.update(k)
        els.append(d)
        return d
    for i, fl in enumerate(niveluri or NIVELURI):
        for k in range(1, ap_per_nivel + 1):
            x0 = 200 + (k - 1) * 150
            add(element_type="contur_apartament", floor=fl, plan_type="ambele", x=x0, y=100,
                label="P%d_%d" % (i, k),
                cable_path=[[x0, 100], [x0 + 140, 100], [x0 + 140, 300], [x0, 300]])
            if cu_interfon and ("P%d_%d" % (i, k)) not in fara_post:
                add(element_type="post_interior_interfon", floor=fl, x=x0 + 20, y=120,
                    mount_height_m=1.5)
        if cu_interfon:
            add(element_type="traseu_cs", floor=fl, x=100, y=200, label="interfon",
                cable_path=[[100, 200], [100, 110], [200 + (ap_per_nivel - 1) * 150 + 20, 110]])
    if cu_interfon:
        if not fara_panou:
            add(element_type="panou_apel_interfon", floor="parter", x=60, y=400, mount_height_m=1.5)
        add(element_type="cititor_control_acces", floor="parter", x=80, y=400, mount_height_m=1.3)
        add(element_type="yala_electromagnetica", floor="parter", x=100, y=400, mount_height_m=2.1)
        if not fara_sursa:
            add(element_type="alimentare_receptor", floor="parter", plan_type="forta", x=70, y=420,
                label="Sursa interfon", power_w=100, phase="mono", mount_height_m=1.5)
    return els


# ── A. MIGRATIA ───────────────────────────────────────────────────────────────────────────────────
def migratii_chk():
    d = os.path.join(RADACINA, "supabase", "migrations")
    out = []
    # DOAR .sql: sonda pe sursa a aratat ca un fisier strain in folder (o copie `.sql.bak`) se sorta
    # DUPA migratie si era citit in locul ei — proba ar fi trecut pe o copie veche.
    for f in sorted(x for x in os.listdir(d) if x.endswith(".sql")):
        t = io.open(os.path.join(d, f), encoding="utf-8").read()
        m = re.search(r"add constraint chk_element_type check \(element_type in \((.*?)\)\);", t, re.S)
        if m:
            out.append((f, re.findall(r"'([a-z_0-9]+)'", re.sub(r"--[^\n]*", "", m.group(1)))))
    return out


def proba_a():
    print("\nA. migratia")
    ms = migratii_chk()
    f, lst = ms[-1]
    v("[A] ultima migratie a constrangerii e a interfonului", "videointerfon" in f, f)
    v("[A] contine cele patru tipuri", TIPURI <= set(lst), sorted(TIPURI - set(lst)))
    v("[A] ADITIVA: nicio valoare din migratia anterioara nu se pierde",
      set(ms[-2][1]) <= set(lst), sorted(set(ms[-2][1]) - set(lst)))
    v("[A] fara dubluri, 73 de valori", len(lst) == len(set(lst)) == 73, "%d / %d" % (len(lst), len(set(lst))))
    import draw_elements as DE
    v("[A] toate tipurile de curenti slabi ale planşei sunt in constrangere",
      set(DE._CS_TYPES) <= set(lst), sorted(set(DE._CS_TYPES) - set(lst)))
    return lst


# ── B. LISTELE DE INREGISTRARE ──────────────────────────────────────────────────────────────────
def _editor_src():
    return io.open(EDITOR, encoding="utf-8").read()


def surse_liste(migratie):
    """{nume lista: multimea tipurilor de interfon pe care le contine}. Fiecare lista e citita din
    SURSA ei (modulul Python incarcat, sau textul TSX), nu dintr-o copie."""
    import draw_elements as DE
    import memoriu_generator as MG
    import caiet_sarcini_generator as CG
    import fitz
    ed = _editor_src()

    def inter(x):
        return TIPURI & set(x)
    # _draw_cs: un tip „e desenat” daca simbolul lui difera de al dozei (ramura implicita) — un tip
    # uitat in functie cade pe doza si ARATA ca un patratel mic, fara nicio eroare.
    def _amprenta(t):
        doc = fitz.open()
        pg = doc.new_page(width=40, height=40)
        DE._draw_cs(pg, 20, 20, t)
        h = hashlib.sha256(pg.get_pixmap(dpi=144).samples).hexdigest()
        doc.close()
        return h
    doza = _amprenta("doza_cs")
    desenate = {t for t in interfon.TIPURI if _amprenta(t) != doza}
    ci = re.search(r"const CS_INTERFON = \[(.*?)\] as const;", ed, re.S)
    ab = re.search(r"const CS_ABBR: Record<string, string> = \{(.*?)\};", ed, re.S)
    cs = ed[ed.index("function csSymbol("):ed.index("const csHit")]
    ct = re.search(r"const CS_TYPES: string\[\] = \[(.*?)\];", ed, re.S)
    return {
        "migratie chk_element_type": inter(migratie),
        "draw_elements._CS_TYPES": inter(DE._CS_TYPES),
        "draw_elements._CS_FAMILY": inter(k for k, c in DE._CS_FAMILY.items() if c == DE._CS_INTERFON),
        "draw_elements._CS_ABBR": inter(k for k, a in DE._CS_ABBR.items() if a),
        "draw_elements._CS_LEGEND": inter(DE._CS_LEGEND),
        "draw_elements._CS_BOM_NAME": inter(DE._CS_BOM_NAME),
        "draw_elements._CS_HEIGHT": inter(DE._CS_HEIGHT),
        "draw_elements._CS_POWER_W": inter(DE._CS_POWER_W),
        "draw_elements._draw_cs (simbol propriu)": desenate,
        "memoriu._CS_NUME": inter(MG._CS_NUME),
        "memoriu._CS_ORDINE": inter(MG._CS_ORDINE),
        "memoriu._CS_INTERFON": inter(MG._CS_INTERFON),
        "caiet._CS_INTERFON": inter(CG._CS_INTERFON),
        "editor CS_INTERFON": inter(re.findall(r'value: "([a-z_]+)"', ci.group(1)) if ci else []),
        "editor CS_TYPES (include familia)": TIPURI if (ct and "CS_INTERFON" in ct.group(1)) else set(),
        "editor CS_ABBR": inter(k for k, a in re.findall(r'(\w+): "([A-Z]*)"', ab.group(1) if ab else "") if a),
        "editor csSymbol": inter(re.findall(r'case "([a-z_]+)":', cs)),
    }


def lipsuri(surse):
    return {nume: sorted(TIPURI - s) for nume, s in surse.items() if TIPURI - s}


def proba_b(migratie):
    print("\nB. listele de inregistrare")
    surse = surse_liste(migratie)
    v("[B] cele %d liste contin toate cele 4 tipuri" % len(surse), not lipsuri(surse), lipsuri(surse))
    # SONDA, pe FIECARE lista: scoti un tip -> proba trebuie sa arate exact lista aia
    tinta = interfon.POST
    prinse = 0
    for nume in surse:
        s2 = copy.deepcopy(surse)
        s2[nume] = s2[nume] - {tinta}
        if list(lipsuri(s2)) == [nume] and lipsuri(s2)[nume] == [tinta]:
            prinse += 1
        else:
            v("[B] sonda: %s fara post -> prins" % nume, False, lipsuri(s2))
    v("[B] sonda: scos pe rand din fiecare lista, prins de %d/%d ori" % (prinse, len(surse)),
      prinse == len(surse))
    # sonda pe SURSA, nu pe copie: se sterge un tip din textul editorului si se reciteste
    ed = _editor_src()
    stricat = ed.replace('case "yala_electromagnetica":', 'case "yala_electromagnetica_X":', 1)
    cs = stricat[stricat.index("function csSymbol("):stricat.index("const csHit")]
    v("[B] sonda pe sursa: un `case` redenumit in csSymbol se vede",
      interfon.YALA not in re.findall(r'case "([a-z_]+)":', cs))

    # VALORILE oglinzilor, nu doar prezenta: inaltimi, abrevieri, culoare, cablu, sursa
    import draw_elements as DE
    import enrich_circuits as EC
    ed = _editor_src()
    ci = re.search(r"const CS_INTERFON = \[(.*?)\] as const;", ed, re.S).group(1)
    h_ed = {t: float(h) for t, h in re.findall(r'value: "([a-z_]+)",\s*label: "[^"]*",\s*h: ([\d.]+)', ci)}
    v("[B] inaltimile editorului = _CS_HEIGHT", h_ed == {t: DE._CS_HEIGHT[t] for t in TIPURI}, h_ed)
    ab = dict(re.findall(r'(\w+): "([A-Z]*)"',
                         re.search(r"const CS_ABBR: Record<string, string> = \{(.*?)\};", ed, re.S).group(1)))
    v("[B] abrevierile editorului = _CS_ABBR", all(ab.get(t) == DE._CS_ABBR[t] for t in TIPURI),
      {t: (ab.get(t), DE._CS_ABBR[t]) for t in TIPURI})
    _abr = [DE._CS_ABBR[t] for t in DE._CS_TYPES if DE._CS_ABBR.get(t)] + list(DE._DET_ABBR.values())
    v("[B] abrevierile noi nu se lovesc de nicio abreviere existenta (CS + detectie)",
      all(_abr.count(DE._CS_ABBR[t]) == 1 for t in TIPURI))
    col = re.search(r'const COL_CS_INT\s*=\s*"#([0-9A-Fa-f]{6})"', ed).group(1)
    v("[B] culoarea editorului = _CS_INTERFON",
      tuple(round(int(col[i:i + 2], 16) / 255, 3) for i in (0, 2, 4)) == DE._CS_INTERFON, col)
    m = re.search(r'\{ value: "interfon",.*?col: COL_CS_INT, dash: \[([\d, ]+)\] \}', ed)
    v("[B] cablul „interfon” in editor, cu linie-punct = backend",
      m is not None and "[%s] 0" % " ".join(x.strip() for x in m.group(1).split(",")) == DE._CS_CABLE["interfon"]["dash"],
      m.group(1) if m else "lipseste")
    v("[B] cablul „interfon” NU inmulteste metrii cu aparatele (magistrala)",
      "interfon" not in DE._CS_CABLE_SERVESTE)
    cst = io.open(CONST, encoding="utf-8").read()
    sb = re.search(r'\{ label: "([^"]+)", default_w: (\d+), default_phase: "mono", default_height: [\d.]+ \}',
                   cst[cst.index("export const BLOC_RECEPTOR_TYPES"):])
    v("[B] eticheta sursei: editor = interfon.SURSA_ETICHETA", sb and sb.group(1) == interfon.SURSA_ETICHETA)
    v("[B] puterea implicita a sursei: editor = enrich",
      sb and int(sb.group(2)) == EC._RECEPTOR_DEFAULT_W[interfon.SURSA_TIP])
    v("[B] eticheta butonului e recunoscuta ca sursa de backend",
      EC.receptor_type_of(interfon.SURSA_ETICHETA) == interfon.SURSA_TIP)
    v("[B] butonul sursei exista, cu poarta de bloc",
      re.search(r'label: "Sursa interfon", btnText: "[^"]+", gate: \{ kind: "bloc" \}', cst) is not None)


# ── C. NUMEROTAREA + SCHEMA, PRIN ENDPOINT ──────────────────────────────────────────────────────
def proba_c():
    print("\nC. P+4 cu interfon: 17 planşe, IE.17 desenata; fara interfon: 16")
    from fastapi.testclient import TestClient
    import main
    import fitz
    c = TestClient(main.app)

    def num(extra):
        r = c.post("/plansa-numbering", json=dict(BLOC_STEAGURI, **extra)).json()
        return r.get("planse") or []
    cu = num({"plan_elements": bloc_proba()})
    fara = num({"plan_elements": bloc_proba(cu_interfon=False)})
    v("[C] cu interfon: 17 planşe", len(cu) == 17, len(cu))
    v("[C] IE.17 = schema_distributie_interfon", cu and cu[-1]["nr"] == "IE.17"
      and cu[-1]["tip"] == "schema_distributie_interfon", cu[-1] if cu else None)
    v("[C] SONDA: fara elementele de interfon -> 16, ultima IE.16 = TCC",
      len(fara) == 16 and fara[-1]["nr"] == "IE.16" and fara[-1]["tip"] == "schema_tcc",
      [(p["nr"], p["tip"]) for p in fara[-2:]])
    v("[C] interfonul nu muta nimic: primele 16 identice cu blocul fara interfon", cu[:16] == fara)
    sw = num({"has_interfon": True})
    v("[C] un `has_interfon: true` trimis de cineva NU aprinde planşa", len(sw) == 16, len(sw))
    necasa = num({"plan_elements": bloc_proba(), "este_bloc": False})
    v("[C] fara `este_bloc` elementele NU se citesc: numerotarea de azi (16)", necasa == fara,
      [(p["nr"], p["tip"]) for p in necasa[-2:]])
    v("[C] `has_interfon` nu mai e in contractul endpointului",
      "has_interfon" not in main.PlansaNumberingRequest.model_fields)

    body = {"plan_elements": bloc_proba(), "plansa_nr": "IE.17",
            "cartus_firma": {"firma_nume": "PROBA SRL"},
            "cartus_proiect": {"titlu_proiect": "Bloc P+4 proba", "faza": "PT"}}
    r = c.post("/generate-schema-interfon-b64", json=body).json()
    v("[C] schema de interfon generata", bool(r.get("pdf_base64")), r.get("error") or r.get("reason"))
    if r.get("pdf_base64"):
        import base64
        d = fitz.open(stream=base64.b64decode(r["pdf_base64"]), filetype="pdf")
        t = d[0].get_text()
        v("[C] desenata: IE.17 in cartus", "IE.17" in t)
        v("[C] titlul = numele din numerotare", "SCHEMA DE DISTRIBUTIE RETEA DE INTERFON" in t)
        _ap = ["P%d_%d" % (i, k) for i in range(5) for k in (1, 2, 3)]
        v("[C] toate cele 15 apartamente, cu eticheta lor", all(a in t for a in _ap),
          [a for a in _ap if a not in t])
        v("[C] toate cele 5 niveluri", all(x in t for x in ("PARTER", "ETAJ 4", "ETAJ 2")))
        v("[C] intrarea: PA, CCA, YE si sursa", all(x in t for x in ("PA", "CCA", "YE", "SURSA VIDEOINTERFON")))
    r2 = c.post("/generate-schema-interfon-b64",
                json=dict(body, plan_elements=bloc_proba(cu_interfon=False))).json()
    v("[C] SONDA: fara elementele de interfon, generatorul sare (aceeasi conditie)",
      r2.get("skipped") is True and not r2.get("pdf_base64"), r2)
    # Schema de curenti slabi NU vede interfonul: un bloc numai cu interfon n-o primeste goala
    import schema_cs
    v("[C] schema de curenti slabi exclude interfonul (bloc numai cu interfon -> fara schema CS)",
      schema_cs.build_cs_schema(bloc_proba()) is None)
    # ...si nici NU O ANUNTA. Cazul real: inginerul regenereaza planşele de curenti slabi ca sa apara
    # interfonul pe ele (has_cs=True). Pana aici numerotarea anunta atunci si schema CS, generatorul o
    # sarea, iar in borderou ramanea o planşa promisa si nelivrata (IE.22 pe blocul de proba).
    cs = num({"plan_elements": bloc_proba(), "has_cs": True})
    tip = [p["tip"] for p in cs]
    v("[C] bloc cu interfon + planşe CS, fara alt echipament: schema CS NEanuntata",
      "schema_cs" not in tip and tip.count("plan_curenti_slabi") == 5
      and tip[-1] == "schema_distributie_interfon", tip[-3:])
    cu_pir = bloc_proba() + [{"id": "el-pir", "element_type": "detector_pir", "floor": "parter",
                              "x": 50, "y": 50, "plan_type": "curenti_slabi"}]
    tip2 = [p["tip"] for p in num({"plan_elements": cu_pir, "has_cs": True})]
    v("[C] SONDA: acelasi bloc cu un PIR -> schema CS anuntata, inaintea interfonului",
      "schema_cs" in tip2 and tip2.index("schema_cs") < tip2.index("schema_distributie_interfon"), tip2[-3:])
    tip3 = [p["tip"] for p in num({"extra_floors": [], "has_tcc": False, "has_distributie": False,
                                   "has_teg": True, "has_cs": True, "este_bloc": False,
                                   "plan_elements": casa_proba()})]
    v("[C] casa cu echipamente CS: schema CS anuntata, ca pana acum", "schema_cs" in tip3, tip3)
    tip4 = [p["tip"] for p in num({"has_cs": True})]
    v("[C] fara elemente (apelantii de azi): schema CS urmeaza planşele, ca pana acum",
      "schema_cs" in tip4, tip4)
    v("[C] generatorul si numerotarea citesc ACEEASI poarta (schema_cs.are_continut)",
      "if not are_continut(elements):" in io.open(os.path.join(RADACINA, "schema_cs.py"), encoding="utf-8").read()
      and "_scs.are_continut(_rows)" in io.open(os.path.join(RADACINA, "main.py"), encoding="utf-8").read())


# ── D. LISTA DE CANTITATI ───────────────────────────────────────────────────────────────────────
def _bom(els):
    import bom
    import draw_elements as DE
    import enrich_circuits as EC
    circs = EC.enrich_circuits(els, {"building_type": "bloc_locuinte"}, base_circuits=[])
    return bom.build_bom(els, circs, DE.compute_cables(els)[0], DE._PX_TO_M)["rows"]


def _rand(rows, den):
    return next((r for r in rows if r["denumire"] == den), None)


def proba_d():
    print("\nD. lista de cantitati")
    rows = _bom(bloc_proba())
    son = _rand(rows, "Buton de sonerie la usa apartamentului")
    ies = _rand(rows, "Buton de iesire pentru deblocarea usii de acces")
    v("[D] 15 posturi -> 15 butoane de sonerie", son and son["cantitate"] == 15, son)
    v("[D] 1 yala -> 1 buton de iesire", ies and ies["cantitate"] == 1, ies)
    els = bloc_proba()
    els.append(dict(next(e for e in els if e["element_type"] == interfon.POST),
                    id="el-9001", x=999, y=999))                      # +1 post (copie a unui post)
    son2 = _rand(_bom(els), "Buton de sonerie la usa apartamentului")
    v("[D] SONDA: +1 post -> +1 buton de sonerie", son2 and son2["cantitate"] == 16, son2)
    els = bloc_proba()
    els.append(dict(next(e for e in els if e["element_type"] == interfon.YALA), id="el-9002", x=5))
    ies2 = _rand(_bom(els), "Buton de iesire pentru deblocarea usii de acces")
    v("[D] SONDA: +1 yala -> +1 buton de iesire", ies2 and ies2["cantitate"] == 2, ies2)
    src = _rand(rows, "Sursa de alimentare videointerfon, magistrala 2 fire")
    v("[D] sursa in sectiunea de curenti slabi", src and src["sectiune"] == "CURENTI SLABI", src)
    v("[D] sursa NU si printre receptoarele de forta (numarata o data)",
      not any(r["categorie"] == "Receptoare" and "interfon" in r["denumire"].lower() for r in rows))
    import draw_elements as DE
    cab = _rand(rows, DE._CS_CABLE["interfon"]["bom"])
    # 5 niveluri x (90 + 420) px desenati, la scara fixa, +10%: lungimea traseului, fara inmultire
    asteptat = round(5 * 510 * DE._PX_TO_M * 1.1, 1)
    v("[D] cablul = lungimea desenata (+10%%), NU inmultita cu posturile (%.1f m)" % asteptat,
      cab and abs(cab["cantitate"] - asteptat) < 0.2 and cab["specificatie"] == "traseu desenat", cab)
    v("[D] bloc fara interfon: niciun rand de interfon",
      not any("interfon" in str(r).lower() or "sonerie" in str(r).lower() for r in _bom(bloc_proba(cu_interfon=False))))


# ── E. SURSA PE TCC ─────────────────────────────────────────────────────────────────────────────
def bloc_cu_comun():
    """Blocul de proba + ce are orice bloc in SPATIUL COMUN: lumina si o priza pe casa scarii,
    pe parter si la etaj. Asa TCC-ul are si alte circuite decat sursa, iar „restul circuitelor TCC
    identice" se poate verifica pe ceva."""
    els = bloc_proba()
    for i, fl in enumerate(("parter", "etaj")):
        for k, (et, x, y, pt) in enumerate((("aplica_tavan", 120, 180, "iluminat"),
                                            ("aplica_tavan", 150, 250, "iluminat"),
                                            ("intrerupator_simplu", 110, 150, "iluminat"),
                                            ("priza_simpla", 170, 280, "forta"))):
            els.append({"id": "com-%d-%d" % (i, k), "project_id": "proba", "floor": fl,
                        "element_type": et, "plan_type": pt, "x": x, "y": y, "room": "Casa scarii",
                        "label": None, "wall_mounted": et != "aplica_tavan", "rotation": 0,
                        "power_w": None, "cable_path": [], "mount_height_m": None, "kit_panica": False})
    return els


# Circuitele unui plan, calculate in arborele dat (HEAD sau cel de lucru) — pentru comparatia de la E.
ENRICH_SCRIPT = r'''
import json, logging, sys
sys.path.insert(0, sys.argv[1])
logging.disable(logging.CRITICAL)
import enrich_circuits as EC
els = json.loads(open(sys.argv[2], encoding="utf-8").read())
sys.stdout.write(json.dumps(EC.enrich_circuits(els, {}, base_circuits=[]), default=str, sort_keys=True))
'''


def proba_e():
    print("\nE. sursa pe TCC, circuit separat, cu protectia ei (decizia lui Dan)")
    import enrich_circuits as EC
    import schema_payloads as SP
    els = bloc_cu_comun()
    circs = EC.enrich_circuits(els, {"building_type": "bloc_locuinte"}, base_circuits=[])
    tcc = [c for c in circs if c.get("panel") == "TCC"]
    src = [c for c in tcc if "Sursa interfon" in str(c.get("description"))]
    v("[E] exact un circuit al sursei, pe TCC", len(src) == 1, [(c["id"], c["description"]) for c in tcc])
    v("[E] TCC are si alte circuite (lumina + priza scarii), nu doar sursa", len(tcc) > 1, len(tcc))
    if src:
        s = src[0]
        v("[E] dedicat, 100 W", s["type"] == "dedicat" and s["power_w"] == 100, (s["type"], s["power_w"]))
        v("[E] sursa: MCB 10 A curba B, cablu CYY-F 3x1,5 (decizia lui Dan)",
          s["breaker_a"] == 10 and str(s["breaker_type"]).endswith("-B") and s["cable_type"] == "CYY-F 3x1.5",
          (s["breaker_a"], s["breaker_type"], s["cable_type"]))
        # SONDA: aceeasi sursa, sub alta eticheta, nu mai e recunoscuta -> regula GENERALA, neatinsa
        gen = copy.deepcopy(els)
        for e in gen:
            if interfon.este_sursa(e):
                e["label"] = "Receptor oarecare"
        g = [c for c in EC.enrich_circuits(gen, {}, base_circuits=[])
             if "Receptor oarecare" in str(c.get("description"))][0]
        v("[E] SONDA: un receptor oarecare de 100 W ramane pe regula generala (16 A, 3x2,5)",
          (g["breaker_a"], g["cable_type"]) == (16, "CYY-F 3x2.5"), (g["breaker_a"], g["cable_type"]))
        sar = SP.sarcini_scheme([{"nr": "IE.16", "tip": "schema_tcc", "nivel": None,
                                  "nume": "SCHEMA ELECTRICĂ MONOFILARĂ TCC"}], circs)["sarcini"]
        v("[E] circuitul sursei e in SCHEMA TCC",
          any("Sursa interfon" in json.dumps(t, ensure_ascii=False) for t in sar), len(sar))
    # RESTUL CIRCUITELOR TCC, fata de HEAD: acelasi plan, calculat in ambii arbori
    lucru = tempfile.mkdtemp(prefix="p7b_e_")
    try:
        f = os.path.join(lucru, "els.json")
        io.open(f, "w", encoding="utf-8").write(json.dumps(els))
        io.open(os.path.join(lucru, "e.py"), "w", encoding="utf-8", newline="").write(ENRICH_SCRIPT)
        rez = {}
        for nume, rad in (("HEAD", head_dir()), ("lucru", RADACINA)):
            r = subprocess.run([sys.executable, os.path.join(lucru, "e.py"), rad, f], capture_output=True)
            rez[nume] = json.loads(r.stdout.decode("utf-8") or "null")
        def _tcc(cc, cu_sursa):
            return [c for c in (cc or []) if c.get("panel") == "TCC"
                    and ("Sursa interfon" in str(c.get("description"))) == cu_sursa]
        v("[E] restul circuitelor TCC: identice cu HEAD (%d circuite)" % len(_tcc(rez["lucru"], False)),
          _tcc(rez["HEAD"], False) == _tcc(rez["lucru"], False) and _tcc(rez["HEAD"], False))
        h_src = _tcc(rez["HEAD"], True)
        v("[E] singura diferenta e sursa: la HEAD 16 A / 3x2,5, acum 10 A / 3x1,5",
          [(c["breaker_a"], c["cable_type"]) for c in h_src] == [(16, "CYY-F 3x2.5")]
          and [(c["breaker_a"], c["cable_type"]) for c in _tcc(rez["lucru"], True)] == [(10, "CYY-F 3x1.5")],
          ([(c["breaker_a"], c["cable_type"]) for c in h_src],))
    finally:
        shutil.rmtree(lucru, ignore_errors=True)
    # SONDA: aceeasi sursa, mutata INTR-UN apartament -> nu mai e comuna, deci nu mai e pe TCC
    mutat = copy.deepcopy(els)
    for e in mutat:
        if interfon.este_sursa(e):
            e["x"], e["y"] = 260, 200                     # in conturul P0_1
    c2 = EC.enrich_circuits(mutat, {}, base_circuits=[])
    v("[E] SONDA: sursa in apartament -> NU pe TCC (regula spatiului comun decide)",
      not any(c.get("panel") == "TCC" and "Sursa interfon" in str(c.get("description")) for c in c2))


# ── F. COPIEREA P4 ──────────────────────────────────────────────────────────────────────────────
def _p4(els, tinta):
    import apartments as AP
    import floors as FL
    return AP.copieri_pentru_nivel(els, tinta, "proba", FL.floor_canonic, FL.floor_index)


def proba_f():
    print("\nF. copierea P4 ia postul interior")
    # doua niveluri cu aceleasi contururi (etichetele: P0_x pe „etaj”, P1_x pe „etaj 2”);
    # sursa „etaj” primeste un post + un bec in fiecare apartament, tinta „etaj 2” doar contururile
    els = bloc_proba(cu_interfon=False, niveluri=["etaj", "etaj 2"], ap_per_nivel=2)
    for e in list(els):
        if e.get("floor") == "etaj" and e["element_type"] == "contur_apartament":
            x0 = e["x"]
            els.append(dict(e, id=e["id"] + "-p", element_type="post_interior_interfon",
                            plan_type="curenti_slabi", x=x0 + 20, y=120, cable_path=[], label=None))
            els.append(dict(e, id=e["id"] + "-b", element_type="aplica_tavan", plan_type="iluminat",
                            x=x0 + 70, y=200, cable_path=[], label=None))
    r = _p4(els, "etaj 2")
    copiate = [el for cp in r["copieri"] for el in cp["elemente"]]
    v("[F] ambele apartamente copiate", len(r["copieri"]) == 2, r.get("sarite"))
    v("[F] postul interior e copiat in fiecare",
      sum(1 for el in copiate if el["element_type"] == interfon.POST) == 2,
      [el["element_type"] for el in copiate])
    v("[F] copiat pe nivelul tinta, pe planşa lui", all(el["floor"] == "etaj 2" and el["plan_type"] ==
                                                          "curenti_slabi" for el in copiate
                                                          if el["element_type"] == interfon.POST))
    # SONDA (golul pe care-l acopera G): un apartament deja copiat nu mai primeste nimic
    for e in els:
        if e.get("floor") == "etaj 2" and e["element_type"] == "contur_apartament" and e["label"] == "P1_1":
            e["copiat_din"] = "altceva"
    r2 = _p4(els, "etaj 2")
    v("[F] SONDA: apartamentul deja copiat e sarit (copierea se face o SINGURA data)",
      [c["tinta"] for c in r2["copieri"]] == ["P1_2"], r2.get("sarite"))


# ── G. AVERTISMENTELE ──────────────────────────────────────────────────────────────────────────
def proba_g():
    print("\nG. avertismentele")
    v("[G] bloc complet: niciun avertisment", interfon.avertismente(bloc_proba()) == [],
      interfon.avertismente(bloc_proba()))
    a = interfon.avertismente(bloc_proba(fara_post=("P2_3",)))
    v("[G] un apartament fara post -> numit", a == ["Apartamentul P2_3 nu are post de interfon"], a)
    a = interfon.avertismente(bloc_proba(fara_post=("P0_1", "P4_2")))
    v("[G] doua -> „Apartamentele P0_1 și P4_2 nu au post de interfon”",
      a == ["Apartamentele P0_1 și P4_2 nu au post de interfon"], a)
    a = interfon.avertismente(bloc_proba(fara_panou=True))
    v("[G] posturi fara panou de apel", "Există posturi de interfon, dar niciun panou de apel" in a, a)
    a = interfon.avertismente(bloc_proba(fara_sursa=True))
    v("[G] sistem fara sursa", any("nu are sursă de alimentare" in x for x in a), a)
    v("[G] bloc FARA interfon: niciun avertisment (neschimbat)",
      interfon.avertismente(bloc_proba(cu_interfon=False)) == [])
    # un post pus pe ALT nivel nu acopera apartamentul (apartenenta e pe nivelul elementului)
    els = bloc_proba(fara_post=("P1_1",))
    els.append({"id": "el-mut", "element_type": interfon.POST, "floor": "etaj 2", "x": 220, "y": 120})
    a = interfon.avertismente(els)
    v("[G] postul de pe alt nivel nu acopera apartamentul", a == ["Apartamentul P1_1 nu are post de interfon"], a)
    # /regenerate-plan le intoarce pe planşa de curenti slabi (citire pe tot proiectul)
    src = io.open(os.path.join(RADACINA, "main.py"), encoding="utf-8").read()
    i = src.index('@app.post("/regenerate-plan")')
    corp = src[i:src.index("\n@app.", i + 10)]
    v("[G] /regenerate-plan pune `avertismente` in raspuns, doar la curenti slabi",
      '_res["avertismente"] = _ifn.avertismente(_toate)' in corp and '"curenti_slabi"' in corp)
    ed = _editor_src()
    v("[G] editorul le afiseaza, sub rezultatul „Obtine plan”",
      "setCsAvertismente(" in ed and 'mode === "curenti_slabi" && csAvertismente.length > 0' in ed)


# ── H. CASA: paleta + documentele, fata de HEAD ─────────────────────────────────────────────────
HAM_TS = r'''
const ts = require("typescript"), fs = require("fs"), path = require("path");
const tr = (src, out, repl) => {
  let js = ts.transpileModule(fs.readFileSync(src, "utf8"),
    { compilerOptions: { module: ts.ModuleKind.ES2020, target: ts.ScriptTarget.ES2020 } }).outputText;
  for (const [a, b] of repl) js = js.split(a).join(b);
  fs.writeFileSync(path.join(__dirname, out), js);
};
tr("lib/floors.ts", "floors.mjs", []);
tr("lib/constants.ts", "constants.mjs", [['"./floors"', '"./floors.mjs"']]);
'''
RUN_TS = r'''
const C = await import("./constants.mjs");
const tipuri = ["casa_unifamiliala", "duplex", "bloc_locuinte", "spatiu_comercial_bloc", ""];
const out = {};
for (const t of tipuri) {
  out[t] = { bloc: C.esteBloc(t),
             butoane: C.visibleEquipmentReceptors({ heatingType: "", enabledEquipment: [], buildingType: t })
                        .map(b => b.label) };
}
out.__fara_tip = C.visibleEquipmentReceptors({ heatingType: "", enabledEquipment: [] }).map(b => b.label);
process.stdout.write(JSON.stringify(out));
'''


def ruleaza_ts():
    d = tempfile.mkdtemp(prefix=".proba-interfon-", dir=APP)
    try:
        io.open(os.path.join(d, "build.cjs"), "w", encoding="utf-8", newline="").write(HAM_TS)
        io.open(os.path.join(d, "run.mjs"), "w", encoding="utf-8", newline="").write(RUN_TS)
        b = subprocess.run(["node", os.path.join(d, "build.cjs")], capture_output=True, text=True, cwd=APP)
        if b.returncode:
            return None, b.stderr[-400:]
        r = subprocess.run(["node", os.path.join(d, "run.mjs")], capture_output=True, cwd=APP)
        return json.loads(r.stdout.decode("utf-8") or "null"), r.stderr.decode("utf-8", "replace")[-400:]
    finally:
        shutil.rmtree(d, ignore_errors=True)


def poarta_editor(ed):
    """Butoanele familiei stau DOAR sub `bloc &&`, iar cablul „interfon” e filtrat pe acelasi `bloc`,
    care vine din `esteBloc(buildingType)` — o singura conditie."""
    i = ed.index("const renderCsSection = () => {")
    corp = ed[i:ed.index("const renderSigurantaSection", i)]
    j = corp.find("grup(CS_INTERFON)")
    return ("const bloc = esteBloc(buildingType);" in corp
            and j > 0 and corp.rfind("{bloc && (", 0, j) > 0
            and corp.count("grup(CS_INTERFON)") == 1
            and 'CS_CABLES.filter(c => bloc || c.value !== "interfon")' in corp)


# Documentele unui proiect, randate de ACELASI script in ARBORELE DE LA HEAD si in cel de lucru:
# circuite, lista de cantitati, memoriu, caiet (randurile DOCX), numerotarea PRIN ENDPOINTUL
# arborelui, schema de curenti slabi si planşele fiecarui nivel (PIXELII, nu octetii PDF — metadatele
# poarta un ceas). Esecul unei randari intra in amprenta ca eroare, deci nu se poate ascunde.
DOC_SCRIPT = r'''
import base64, hashlib, io, json, logging, sys
sys.path.insert(0, sys.argv[1])
logging.disable(logging.CRITICAL)
import fitz
import draw_elements as DE, enrich_circuits as EC, bom, memoriu_generator as MG
import caiet_sarcini_generator as CG, schema_cs
from docx import Document
cfg = json.loads(open(sys.argv[2], encoding="utf-8").read())
els = cfg["els"]
out = {}
circs = EC.enrich_circuits(els, {"building_type": cfg["tip"]}, base_circuits=[])
out["circuite"] = circs
out["bom"] = bom.build_bom(els, circs, DE.compute_cables(els)[0], DE._PX_TO_M)["rows"]
data = {"cartus_proiect": {"faza": "DTAC+PT", "titlu_proiect": "Proba"}, "cartus_firma": {},
        "circuits": circs, "planse": [], "has_cs": True, "extra_floors": cfg["extra_floors"]}
out["memoriu"] = [p.text for p in Document(io.BytesIO(MG.build_memoriu_docx(dict(data)))).paragraphs]
out["caiet"] = [p.text for p in Document(io.BytesIO(CG.build_caiet_docx(dict(data)))).paragraphs]
from fastapi.testclient import TestClient
import main
out["numerotare"] = TestClient(main.app).post("/plansa-numbering", json=cfg["numerotare"]).json()
def px(pdf):
    if not pdf:
        return None
    d = fitz.open(stream=pdf, filetype="pdf")
    h = hashlib.sha256()
    for pg in d:
        h.update(pg.get_pixmap(dpi=90).samples)
    return h.hexdigest()
out["schema_cs_px"] = px(schema_cs.build_cs_schema(els, {}, {"faza": "PT"}, "IE.3"))
baza = fitz.open(); baza.new_page(width=842, height=595)
b64 = base64.b64encode(baza.tobytes()).decode()
for fl in sorted({e.get("floor") or "parter" for e in els}):
    for tip in ("curenti_slabi", "iluminat", "forta"):
        r = DE.redraw_from_plan_elements(b64, [e for e in json.loads(json.dumps(els))
                                               if (e.get("floor") or "parter") == fl],
                                         draw_plan_type=tip, rooms=[])
        out["plansa_%s_%s" % (fl, tip)] = px(base64.b64decode(r["pdf_base64"])) if r.get("success") else r
sys.stdout.write(json.dumps(out, ensure_ascii=False, sort_keys=True, default=str))
'''


def casa_proba():
    els, n = [], [0]

    def add(**k):
        n[0] += 1
        d = {"id": "c-%03d" % n[0], "project_id": "casa", "floor": "parter", "label": None,
             "room": "Living", "wall_mounted": False, "rotation": 0, "power_w": None,
             "cable_path": [], "mount_height_m": None, "kit_panica": False, "plan_type": "iluminat"}
        d.update(k)
        els.append(d)
    add(element_type="tablou_teg", plan_type="ambele", x=80, y=80)
    add(element_type="lustra_led", x=300, y=200)
    add(element_type="intrerupator_simplu", x=250, y=150, wall_mounted=True)
    add(element_type="priza_dubla", plan_type="forta", x=320, y=260, wall_mounted=True)
    add(element_type="alimentare_receptor", plan_type="forta", x=400, y=260, label="Aer conditionat")
    for et, x in (("centrala_efractie", 120), ("tastatura_efractie", 140), ("detector_pir", 360),
                  ("camera_video", 420), ("nvr", 160), ("rack_9u", 180), ("priza_date", 330),
                  ("priza_tv", 340), ("sirena_exterioara", 500)):
        add(element_type=et, plan_type="curenti_slabi", x=x, y=120, wall_mounted=True,
            label="interior" if et == "camera_video" else None,
            camera_tip="dome" if et == "camera_video" else None)
    add(element_type="receptor_internet", plan_type="forta", x=180, y=140, label="internet")
    add(element_type="traseu_cs", plan_type="curenti_slabi", x=180, y=120, label="utp",
        cable_path=[[180, 120], [420, 120]])
    return els


def proba_h():
    print("\nH. casa: paleta si documentele")
    rez, err = ruleaza_ts()
    v("[H] functiile reale din constants.ts au rulat", rez is not None, err)
    if rez:
        for t in ("casa_unifamiliala", "duplex", "spatiu_comercial_bloc"):
            v("[H] %-22s nu e bloc, n-are butonul sursei" % t,
              rez[t]["bloc"] is False and "Sursa interfon" not in rez[t]["butoane"], rez[t])
            v("[H] %-22s paleta de receptoare = cea fara tip (neschimbata)" % t,
              rez[t]["butoane"] == rez["__fara_tip"], rez[t]["butoane"])
        v("[H] bloc_locuinte: e bloc si are butonul sursei",
          rez["bloc_locuinte"]["bloc"] is True and "Sursa interfon" in rez["bloc_locuinte"]["butoane"])
    ed = _editor_src()
    v("[H] editorul: familia si cablul ei stau sub `esteBloc`", poarta_editor(ed))
    v("[H] SONDA: poarta scoasa din editor -> prinsa",
      not poarta_editor(ed.replace("{bloc && (", "{true && (", 1)))
    v("[H] SONDA: filtrul cablului scos -> prins",
      not poarta_editor(ed.replace('CS_CABLES.filter(c => bloc || c.value !== "interfon")', "CS_CABLES", 1)))

    # DOCUMENTELE, fata de ARBORELE DE LA HEAD: aceeasi casa / acelasi bloc fara interfon, acelasi
    # script, randat in ambii arbori. Nu discul cu el insusi — cerinta pachetului.
    lucru = tempfile.mkdtemp(prefix="p7b_doc_")
    try:
        head = head_dir()
        io.open(os.path.join(lucru, "doc.py"), "w", encoding="utf-8", newline="").write(DOC_SCRIPT)

        def randeaza(rad, cfg, nume):
            f = os.path.join(lucru, nume + ".json")
            io.open(f, "w", encoding="utf-8").write(json.dumps(cfg))
            r = subprocess.run([sys.executable, os.path.join(lucru, "doc.py"), rad, f],
                               capture_output=True, cwd=lucru)
            return (json.loads(r.stdout.decode("utf-8") or "null") if r.returncode == 0
                    else r.stderr.decode("utf-8", "replace")[-600:])
        bloc_fara = bloc_proba(cu_interfon=False)
        cazuri = {
            "casa": {"els": casa_proba(), "tip": "casa_unifamiliala", "extra_floors": [],
                     "numerotare": {"extra_floors": [], "has_cs": True, "plan_elements": casa_proba()}},
            "bloc fara interfon": {"els": bloc_fara, "tip": "bloc_locuinte",
                                   "extra_floors": BLOC_STEAGURI["extra_floors"],
                                   "numerotare": dict(BLOC_STEAGURI, plan_elements=bloc_fara)},
        }
        for caz, cfg in cazuri.items():
            a, b = randeaza(head, cfg, "h"), randeaza(RADACINA, cfg, "w")
            v("[H] %s: documentele s-au generat in ambii arbori" % caz,
              isinstance(a, dict) and isinstance(b, dict), a if not isinstance(a, dict) else b)
            if not (isinstance(a, dict) and isinstance(b, dict)):
                continue
            erori = sorted(k for k in b if isinstance(b[k], dict) and b[k].get("success") is False)
            v("[H] %s: nicio randare esuata (altfel „identic” n-ar spune nimic)" % caz, not erori, erori)
            dif = sorted(k for k in set(a) | set(b) if a.get(k) != b.get(k))
            v("[H] %s: %d amprente, toate identice cu HEAD" % (caz, len(b)), not dif, dif)
            if caz == "bloc fara interfon":
                pn = b["numerotare"].get("planse") or []
                v("[H] bloc fara interfon: 16 planşe, ultima IE.16 (si la HEAD)",
                  len(pn) == 16 and pn[-1]["nr"] == "IE.16" and a["numerotare"] == b["numerotare"],
                  [p["nr"] for p in pn[-1:]])
            if caz == "casa":
                # SONDA: acelasi arbore, dar cu un post de interfon pe casa -> proba TREBUIE sa vada
                cu = dict(cfg, els=cfg["els"] + [{
                    "id": "c-int", "project_id": "casa", "floor": "parter", "room": "Hol",
                    "element_type": interfon.POST, "plan_type": "curenti_slabi", "x": 200, "y": 300,
                    "label": None, "wall_mounted": True, "rotation": 0, "power_w": None,
                    "cable_path": [], "mount_height_m": 1.5, "kit_panica": False}])
                c = randeaza(RADACINA, cu, "c")
                dif = sorted(k for k in b if isinstance(c, dict) and c.get(k) != b[k])
                v("[H] SONDA: aceeasi casa cu un post de interfon -> proba VEDE diferenta",
                  {"bom", "memoriu", "caiet", "plansa_parter_curenti_slabi"} <= set(dif), dif)
    finally:
        shutil.rmtree(lucru, ignore_errors=True)


# ── J. CITIREA DIN BAZA LA NUMEROTARE: DOAR LA BLOC (ajustarea 1) ───────────────────────────────
class _BazaCazuta:
    """Baza indisponibila: orice interogare arunca. Numara incercarile, ca proba sa vada si DACA
    s-a incercat citirea, nu doar ca raspunsul e bun."""
    def __init__(self):
        self.apeluri = 0

    def table(self, _t):
        self.apeluri += 1
        raise RuntimeError("STUB: baza indisponibila")


def _cu_baza_cazuta(fn):
    import types
    vechi = sys.modules.get("supabase_client")
    baza = _BazaCazuta()
    mod = types.ModuleType("supabase_client")
    mod.supabase = baza
    sys.modules["supabase_client"] = mod
    try:
        return fn(), baza.apeluri
    finally:
        if vechi is not None:
            sys.modules["supabase_client"] = vechi
        else:
            sys.modules.pop("supabase_client", None)


# Cererile EXACT cum le trimite nodul n8n (dupa patch): project_id + este_bloc de la /api/finalize.
CERERE_CASA = {"extra_floors": ["etaj 1"], "has_cs": True, "project_id": "p-casa", "este_bloc": False}
CERERE_BLOC = dict(BLOC_STEAGURI, project_id="p-bloc")


def _casa_merge(c):
    import plansa_numbering as PN
    r, n = _cu_baza_cazuta(lambda: c.post("/plansa-numbering", json=CERERE_CASA).json())
    azi = PN.compute_plansa_numbering(["etaj 1"], False, has_cs=True)   # numerotarea de azi
    return r.get("success") is True and r.get("planse") == azi and n == 0, (r.get("success"), n, r.get("error"))


# Handler-ul REAL al /api/finalize (ca in test_poarta_tipuri): transpilat, legat de stub-uri. Proba
# citeste corpul trimis la n8n si verifica `este_bloc` pentru fiecare tip de cladire.
HAM_FIN = {
    "build.cjs": r'''
const ts = require("typescript"), fs = require("fs"), path = require("path");
const tr = (src, out, repl, pre) => {
  let text = fs.readFileSync(src, "utf8");
  for (const [a, b] of (pre || [])) { if (!text.includes(a)) throw new Error("ancora sondei: " + a); text = text.split(a).join(b); }
  let js = ts.transpileModule(text, { compilerOptions: { module: ts.ModuleKind.ES2020, target: ts.ScriptTarget.ES2020 } }).outputText;
  js = js.replace(/import \{([^}]+)\} from "next\/server";/,
                  (_m, nume) => `import __ns from "next/server.js"; const {${nume}} = __ns;`);
  for (const [a, b] of repl) js = js.split(a).join(b);
  fs.writeFileSync(path.join(__dirname, out), js);
};
tr("lib/floors.ts", "floors.mjs", []);
tr("lib/constants.ts", "constants.mjs", [['"./floors"', '"./floors.mjs"']]);
const imp = [['"next/headers"', '"./stubs/headers.mjs"'], ['"@/lib/supabase"', '"./stubs/supabase.mjs"'],
  ['"@/lib/constants"', '"./constants.mjs"'], ['"@/lib/registru-finalizari"', '"./stubs/registru.mjs"'],
  ['"@/lib/supabaseAdmin"', '"./stubs/admin.mjs"'], ['"@/lib/floors"', '"./floors.mjs"']];
tr("app/api/finalize/route.ts", "route.mjs", imp);
// SONDA: aceeasi ruta, fara conditie — orice proiect ar pleca spre numerotare ca bloc
tr("app/api/finalize/route.ts", "route_sonda.mjs", imp, [['esteBloc(String(inputData.building_type || ""))', 'true']]);
''',
    "run.mjs": r'''
const corpuri = [];
globalThis.fetch = async (u, o) => {
  if (String(u).includes("zynapse-finalize")) {
    corpuri.push(JSON.parse(o.body));
    return { ok: true, status: 200, text: async () => '{"success": false, "error": "stub n8n"}',
             json: async () => ({ success: false }) };
  }
  return { ok: false, status: 503, json: async () => ({ success: false }), text: async () => "{}" };
};
const { NextRequest } = (await import("next/server.js")).default;
const out = {};
for (const mod of ["route.mjs", "route_sonda.mjs"]) {
  const { POST } = await import("./" + mod);
  out[mod] = {};
  for (const tip of ["casa_unifamiliala", "duplex", "spatiu_comercial_bloc", "bloc_locuinte"]) {
    globalThis.__TIP = tip; corpuri.length = 0;
    const req = new NextRequest("http://localhost/api/finalize",
                                { method: "POST", body: JSON.stringify({ project_id: "p1" }) });
    try { await POST(req); } catch (e) { out[mod][tip + "_exc"] = String(e.message || e); }
    out[mod][tip] = corpuri.length ? corpuri[0].este_bloc : "NETRIMIS";
  }
}
process.stdout.write(JSON.stringify(out));
''',
    "stubs/headers.mjs": "export async function cookies() { return { get: () => undefined }; }\n",
    "stubs/registru.mjs": ("export const culegeDocumente = () => [];\n"
                           "export const inregistreaza = async () => ({ ok: false });\n"),
    "stubs/admin.mjs": 'export function createAdminClient() { throw new Error("STUB: admin"); }\n',
    # Proiectul: un circuit in result_data (ca finalizarea sa aiba ce trimite) si TIPUL cladirii din
    # globalThis.__TIP. Restul bazei intoarce gol; backendul (fetch) e „indisponibil".
    "stubs/supabase.mjs": r'''
const proiect = () => ({ data: { result_data: { circuits: [{ id: "C1", panel: "TEG", type: "iluminat",
  description: "Iluminat", floor: 0 }], power_summary: {} }, faza: "DTAC", phase: "DTAC",
  input_data: { building_type: globalThis.__TIP } }, error: null });
function lant(t) {
  return new Proxy(function () {}, {
    get: (_x, k) => k === "then" ? undefined
      : (k === "single" || k === "maybeSingle")
        ? async () => (t === "projects" ? proiect() : { data: {}, error: null })
        : lant(t),
    apply: () => lant(t),
  });
}
export function createServerClient() {
  return { auth: { getUser: async () => ({ data: { user: { id: "u-proba" } }, error: null }) },
           from: (t) => lant(t), rpc: async () => ({ data: null, error: null }),
           storage: { from: () => lant("storage") } };
}
''',
}


def ruleaza_finalize():
    d = tempfile.mkdtemp(prefix=".proba-finalize-", dir=APP)
    try:
        for nume, cod in HAM_FIN.items():
            cale = os.path.join(d, nume)
            os.makedirs(os.path.dirname(cale), exist_ok=True)
            io.open(cale, "w", encoding="utf-8", newline="").write(cod)
        b = subprocess.run(["node", os.path.join(d, "build.cjs")], capture_output=True, text=True, cwd=APP)
        if b.returncode:
            return None, (b.stdout + b.stderr)[-500:]
        r = subprocess.run(["node", os.path.join(d, "run.mjs")], capture_output=True, cwd=APP)
        out = r.stdout.decode("utf-8", "replace")
        return (json.loads(out) if out.strip() else None), r.stderr.decode("utf-8", "replace")[-500:]
    finally:
        shutil.rmtree(d, ignore_errors=True)


def proba_j():
    print("\nJ. citirea din baza la numerotare: doar la bloc")
    from fastapi.testclient import TestClient
    import inspect
    import main
    c = TestClient(main.app)
    ok, det = _casa_merge(c)
    v("[J] casa, baza CAZUTA: numerotarea de azi, fara nicio incercare de citire", ok, det)
    r, n = _cu_baza_cazuta(lambda: c.post("/plansa-numbering", json=CERERE_BLOC).json())
    v("[J] bloc, baza CAZUTA: eroare explicita (nodul n8n reincearca, apoi opreste)",
      r.get("success") is False and "baza indisponibila" in str(r.get("error")) and n >= 1,
      (r.get("success"), n, str(r.get("error"))[:60]))
    # SONDA pe sursa: poarta scoasa din `_elemente_numerotare` -> proba casei TREBUIE sa pice
    orig = main._elemente_numerotare
    src = inspect.getsource(orig)
    mut = src.replace("    if not request.este_bloc:\n        return None\n", "")
    try:
        exec(mut, main.__dict__)
        ok2, det2 = _casa_merge(c)
        v("[J] SONDA: fara conditie, proba casei pica (citeste si cade pe baza)",
          mut != src and not ok2, det2)
    finally:
        main._elemente_numerotare = orig
    ok3, _ = _casa_merge(c)
    v("[J] dupa sonda, functia originala e la loc", ok3)

    # /api/finalize: conditia UNICA (`esteBloc`) decide steagul trimis spre numerotare
    rez, err = ruleaza_finalize()
    v("[J] handler-ul real al /api/finalize a rulat", rez is not None, err)
    if rez:
        r0 = rez["route.mjs"]
        v("[J] finalize: casa -> este_bloc false", r0.get("casa_unifamiliala") is False, r0)
        v("[J] finalize: duplex -> este_bloc false", r0.get("duplex") is False, r0)
        v("[J] finalize: spatiu comercial -> este_bloc false", r0.get("spatiu_comercial_bloc") is False, r0)
        v("[J] finalize: bloc_locuinte -> este_bloc true", r0.get("bloc_locuinte") is True, r0)
        v("[J] SONDA: ruta fara `esteBloc` trimite casa ca bloc -> prinsa",
          rez["route_sonda.mjs"].get("casa_unifamiliala") is True, rez["route_sonda.mjs"])


def proba_i():
    print("\nI. tsc")
    r = subprocess.run(["npx", "tsc", "--noEmit", "-p", "."], capture_output=True, text=True, cwd=APP,
                       shell=(os.name == "nt"))
    v("[I] tsc --noEmit curat", r.returncode == 0 and not r.stdout.strip(), (r.stdout + r.stderr)[-600:])


def main():
    # `python test_interfon.py B D` ruleaza doar sectiunile cerute (A ruleaza mereu: B are nevoie de
    # lista migratiei). Fara argumente — toate.
    cerute = {a.upper() for a in sys.argv[1:]} or set("ABCDEFGHIJ")
    try:
        mig = proba_a()
        for lit, fn in (("B", lambda: proba_b(mig)), ("C", proba_c), ("D", proba_d), ("E", proba_e),
                        ("F", proba_f), ("G", proba_g), ("H", proba_h), ("J", proba_j), ("I", proba_i)):
            if lit in cerute:
                fn()
    finally:
        if _HEAD:
            shutil.rmtree(_HEAD["tmp"], ignore_errors=True)
    print()
    if rele:
        print("ESUAT (%d): %s" % (len(rele), "; ".join(rele)))
        return 1
    print("OK — videointerfonul e inregistrat peste tot, iar planşa se aprinde doar din elemente")
    return 0


if __name__ == "__main__":
    sys.exit(main())
