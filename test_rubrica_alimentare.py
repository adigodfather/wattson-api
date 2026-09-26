"""Rubrica 'Sursa de alimentare': doar la spatiu comercial in bloc (si la un proiect reluat cu
'din firida' salvata).

DE CE EXISTA. Rubrica aparea la TOATE tipurile de cladire, desi are sens doar intr-un imobil
existent. Acum afisarea o decide o singura functie pura, `arataRubricaAlimentare`.

Proba pazeste doua lucruri care arata bine si pot fi gresite:
  1. Ascunderea nu schimba ce PLEACA la generare. O casa trebuie sa trimita tot
     'bransament_propriu' — linia de trimitere nu se atinge.
  2. Nu exista regula pe server, fiindca in aval "" si 'bransament_propriu' dau ACELASI document.
     Asta nu e o presupunere: se genereaza memoriul (cu breviarul), caietul si BOM-ul pe proiecte
     REALE, cu cele trei valori. SONDA DE CONTROL: 'din_firida' TREBUIE sa difere — altfel proba
     nu vede campul, si 'identic' n-ar dovedi nimic.

Rulare:  python test_rubrica_alimentare.py
"""
import copy
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.request

RADACINA = os.path.dirname(os.path.abspath(__file__))
APP = os.path.join(RADACINA, "apps", "zynapse-configurator")
rele = []


def v(nume, cond, det=""):
    print("  %-66s %s %s" % (nume[:66], "OK" if cond else "**ESUAT**", det if not cond else ""))
    if not cond:
        rele.append(nume)


def bloc(sursa, antet):
    i = sursa.index(antet)
    linii = sursa[i:].splitlines()
    out = [linii[0]]
    for l in linii[1:]:
        out.append(l)
        if l == "}":
            break
    return "\n".join(out)


def ruleaza_ts():
    """Compileaza CHIAR functiile din constants.ts (decupate prin text) si le cheama din Node."""
    src = io.open(os.path.join(APP, "lib", "constants.ts"), encoding="utf-8").read()
    bucati = [bloc(src, "export function defaultAlimentare("),
              bloc(src, "export function arataRubricaAlimentare(")]
    lucru = tempfile.mkdtemp(prefix="alim_")
    try:
        f = os.path.join(lucru, "a.ts")
        io.open(f, "w", encoding="utf-8").write("\n\n".join(bucati) + "\n")
        subprocess.run(["npx", "tsc", f, "--outDir", lucru, "--target", "es2020", "--module", "es2020"],
                       capture_output=True, text=True, shell=(os.name == "nt"), cwd=APP)
        shutil.move(os.path.join(lucru, "a.js"), os.path.join(lucru, "a.mjs"))
        io.open(os.path.join(lucru, "d.mjs"), "w", encoding="utf-8").write("""
import { defaultAlimentare, arataRubricaAlimentare } from "./a.mjs";
const TIPURI = ["casa_unifamiliala","duplex","bloc_locuinte","hotel_pensiune","spatiu_comercial_bloc",
  "camin_cultural","scoala","spital","institutie","biserica","sala_sport",
  "hala_productie","depozit","atelier","ferma","statie_tehnologica","tip_necunoscut","",null,undefined];
const out = { tabel: {}, cuFirida: {}, payload: {} };
for (const t of TIPURI) {
  out.tabel[String(t)] = arataRubricaAlimentare(t, "");
  out.cuFirida[String(t)] = arataRubricaAlimentare(t, "din_firida");
}
// exact expresia din configurator: form.alimentare || defaultAlimentare(form.building_type)
const pl = (bt, a) => a || defaultAlimentare(bt);
out.payload = {
  casa_gol: pl("casa_unifamiliala", ""), duplex_gol: pl("duplex", ""),
  comercial_gol: pl("spatiu_comercial_bloc", ""),
  comercial_bransament: pl("spatiu_comercial_bloc", "bransament_propriu"),
  comercial_firida: pl("spatiu_comercial_bloc", "din_firida"),
};
out.reluare = {
  casa_cu_firida: arataRubricaAlimentare("casa_unifamiliala", "din_firida"),
  casa_cu_bransament: arataRubricaAlimentare("casa_unifamiliala", "bransament_propriu"),
  casa_fara_cheie: arataRubricaAlimentare("casa_unifamiliala", undefined),
};
process.stdout.write(JSON.stringify(out));
""")
        r = subprocess.run(["node", os.path.join(lucru, "d.mjs")], capture_output=True, text=True,
                           shell=(os.name == "nt"))
        return json.loads(r.stdout) if r.stdout.strip() else None
    finally:
        shutil.rmtree(lucru, ignore_errors=True)


def text_docx(b):
    """Continutul documentului, pe RANDURI (paragrafe + celule). NU octetii: docx-ul poarta in
    metadate data crearii, deci doi octeti diferiti n-ar spune nimic despre continut."""
    import docx
    d = docx.Document(io.BytesIO(b))
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text for c in row.cells))
    return out


def main():
    # ── A–D: logica de afisare si ce pleaca la generare ─────────────────────────────────────
    r = ruleaza_ts()
    v("[0] functiile din constants.ts s-au putut rula", r is not None)
    if r:
        t = r["tabel"]
        for tip in ("casa_unifamiliala", "duplex", "bloc_locuinte", "hotel_pensiune", "scoala",
                    "spital", "hala_productie", "depozit", "tip_necunoscut", "", "null", "undefined"):
            v("[A] %-22s -> rubrica ASCUNSA" % (tip or '""'), t[tip] is False)
        v("[A] spatiu_comercial_bloc -> rubrica APARE", t["spatiu_comercial_bloc"] is True)
        v("[A] ORICE tip + 'din_firida' -> rubrica apare",
          all(r["cuFirida"].values()), str([k for k, x in r["cuFirida"].items() if not x]))
        p = r["payload"]
        v("[B] casa, rubrica ascunsa -> pleaca 'bransament_propriu'", p["casa_gol"] == "bransament_propriu")
        v("[B] duplex, rubrica ascunsa -> pleaca 'bransament_propriu'", p["duplex_gol"] == "bransament_propriu")
        v("[C] spatiu comercial, implicit -> 'din_firida'", p["comercial_gol"] == "din_firida")
        v("[C] spatiu comercial: AMBELE optiuni ajung in payload",
          p["comercial_bransament"] == "bransament_propriu" and p["comercial_firida"] == "din_firida")
        rl = r["reluare"]
        v("[D] casa reluata cu 'din_firida' salvata -> rubrica APARE", rl["casa_cu_firida"] is True)
        v("[D] casa reluata cu 'bransament_propriu' -> rubrica ascunsa", rl["casa_cu_bransament"] is False)
        v("[D] casa veche, fara cheie in input_data -> rubrica ascunsa", rl["casa_fara_cheie"] is False)

    cfg = io.open(os.path.join(APP, "components", "configurator.tsx"), encoding="utf-8").read()
    v("[B] linia de trimitere e NEATINSA",
      "alimentare: form.alimentare || defaultAlimentare(form.building_type)," in cfg)
    v("[*] configuratorul CHEAMA functia, nu rescrie conditia",
      "arataRubricaAlimentare(form.building_type, form.alimentare) && (" in cfg)
    v("[*] rubrica nu mai apare necondiționat",
      cfg.count('<SelectField label="Sursa de alimentare"') == 1
      and cfg.index("arataRubricaAlimentare(form.building_type") < cfg.index('<SelectField label="Sursa de alimentare"'))

    # ── E: backend, pe proiecte REALE, cu cele trei valori ──────────────────────────────────
    t = io.open(os.path.join(APP, ".env.local"), encoding="utf-8").read() \
        if os.path.exists(os.path.join(APP, ".env.local")) else ""
    env = dict(re.findall(r'^([A-Z0-9_]+)\s*=\s*"?([^"\n\r]*)"?', t, re.M))
    if not env.get("SUPABASE_SERVICE_ROLE_KEY"):
        print("  [E] sarit: fara SUPABASE_SERVICE_ROLE_KEY in .env.local")
    else:
        sys.path.insert(0, RADACINA)
        import memoriu_generator as mg
        import caiet_sarcini_generator as cg
        import bom
        U, K = env["NEXT_PUBLIC_SUPABASE_URL"].rstrip("/"), env["SUPABASE_SERVICE_ROLE_KEY"]
        H = {"apikey": K, "Authorization": "Bearer " + K}
        proj = json.loads(urllib.request.urlopen(urllib.request.Request(
            U + "/rest/v1/projects?select=id,input_data,result_data,project_info,power_summary,circuits,phase"
                "&finalized=eq.true&order=created_at.desc", headers=H), timeout=180).read().decode())
        alese = {}
        for p in proj:
            bt = (p.get("input_data") or {}).get("building_type")
            # power_summary/circuits stau in result_data; coloanele cu acelasi nume sunt goale
            ps_ = p.get("power_summary") or (p.get("result_data") or {}).get("power_summary")
            if bt in ("casa_unifamiliala", "duplex") and bt not in alese and ps_:
                alese[bt] = p
        v("[E] exista cate un proiect real de casa si de duplex", len(alese) == 2, str(list(alese)))
        for bt, p in alese.items():
            circuits = p.get("circuits") or (p.get("result_data") or {}).get("circuits") or []
            ps = p.get("power_summary") or (p.get("result_data") or {}).get("power_summary") or {}
            baza = {"cartus_proiect": {"faza": "DTAC+PT", "nume_proiect": "proba", **(p.get("project_info") or {})},
                    "cartus_firma": {}, "planse": [], "circuits": circuits, "power_summary": ps}

            def doc(gen, alim):
                d = copy.deepcopy(baza); d["alimentare"] = alim
                return text_docx(gen(d))

            def bomrows(alim):
                return json.dumps(bom.build_bom([], circuits, [], 0.01, power_summary=ps, alimentare=alim),
                                  sort_keys=True, default=str)

            eti = "%s %s" % (bt, p["id"][:8])
            m0, m0b = doc(mg.build_memoriu_docx, ""), doc(mg.build_memoriu_docx, "")
            m1, m2 = doc(mg.build_memoriu_docx, "bransament_propriu"), doc(mg.build_memoriu_docx, "din_firida")
            c0, c1, c2 = (doc(cg.build_caiet_docx, a) for a in ("", "bransament_propriu", "din_firida"))
            b0, b1, b2 = (bomrows(a) for a in ("", "bransament_propriu", "din_firida"))

            v("[E] %s: CONTROL — acelasi memoriu generat de doua ori = identic" % eti, m0 == m0b)
            v("[E] %s: memoriu (cu breviar)  \"\" == bransament_propriu" % eti, m0 == m1,
              "%d randuri diferite" % sum(1 for a, b in zip(m0, m1) if a != b))
            v("[E] %s: caiet               \"\" == bransament_propriu" % eti, c0 == c1)
            v("[E] %s: BOM                 \"\" == bransament_propriu" % eti, b0 == b1)
            # ── SONDA DE CONTROL: din_firida TREBUIE sa difere, altfel proba e oarba ──
            v("[E] %s: SONDA memoriu — din_firida DIFERA" % eti, m2 != m0)
            v("[E] %s: SONDA memoriu — apare fraza despre firida" % eti,
              any("firid" in x.lower() for x in m2) and not any("firida blocului" in x.lower() for x in m0))
            v("[E] %s: SONDA caiet — din_firida DIFERA" % eti, c2 != c0)
            v("[E] %s: SONDA BOM — randul de bransament se schimba" % eti,
              "bransament TEG (fix)" in b0 and "alimentare din firida blocului" in b2
              and "bransament TEG (fix)" not in b2)

    print()
    if rele:
        print("ESUAT: " + "; ".join(rele))
        return 1
    print("OK — rubrica apare doar unde are sens; ce pleaca la generare si documentele raman neschimbate")
    return 0


if __name__ == "__main__":
    sys.exit(main())
